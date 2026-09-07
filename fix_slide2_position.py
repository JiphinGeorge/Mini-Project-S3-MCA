import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import subprocess, time

PPTX_PATH = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation.pptx"
prs = pptx.Presentation(PPTX_PATH)

s2 = prs.slides[1] # Slide 2
sp_tree = s2.shapes._spTree

# Identify shapes 6 to 15 (the two top cards, their icons, titles, and body texts)
shapes_to_replace = []
for sh in s2.shapes:
    if sh.top >= Inches(1.6) and sh.top < Inches(4.0):
        shapes_to_replace.append(sh)

print(f"Replacing {len(shapes_to_replace)} top card shapes on Slide 2...")
for sh in shapes_to_replace:
    sp_tree.remove(sh._element)

# Colors
DARK_NAVY = RGBColor(0x0B, 0x1F, 0x3A)
PURPLE_PRIMARY = RGBColor(0x6C, 0x4A, 0xB6)
TEAL_PRIMARY = RGBColor(0x0E, 0x7C, 0x86)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER_GRAY = RGBColor(0xE2, 0xE8, 0xF0)
SLATE_TEXT = RGBColor(0x33, 0x41, 0x55)

# -------------------------------------------------------------
# 1. LEFT CARD: PROBLEM
# -------------------------------------------------------------
card1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.50), Inches(1.70), Inches(5.90), Inches(2.20))
card1.fill.solid()
card1.fill.fore_color.rgb = WHITE
card1.line.color.rgb = BORDER_GRAY
card1.line.width = Pt(1.5)

# Circle Icon for Problem (self-contained with centered text)
icon1 = s2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.85), Inches(1.92), Inches(0.48), Inches(0.48))
icon1.fill.solid()
icon1.fill.fore_color.rgb = PURPLE_PRIMARY
icon1.line.fill.background()
tf1 = icon1.text_frame
tf1.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf1.paragraphs[0]
p.text = "!"
p.font.name = "Calibri"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER

# Title for Problem
t1 = s2.shapes.add_textbox(Inches(1.48), Inches(1.92), Inches(4.70), Inches(0.48))
tf_t1 = t1.text_frame
tf_t1.vertical_anchor = MSO_ANCHOR.MIDDLE
tf_t1.margin_left = Inches(0.05)
p = tf_t1.paragraphs[0]
p.text = "Problem"
p.font.name = "Cambria"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = DARK_NAVY

# Body Text for Problem
b1 = s2.shapes.add_textbox(Inches(0.85), Inches(2.55), Inches(5.25), Inches(1.20))
tf_b1 = b1.text_frame
tf_b1.word_wrap = True
tf_b1.margin_left = Inches(0.0)
tf_b1.margin_right = Inches(0.0)
tf_b1.margin_top = Inches(0.0)
p = tf_b1.paragraphs[0]
p.text = "Clinical transcription data is unstructured and can belong to any of several medical specialties, making manual categorization by medical coders slow, inconsistent, and prone to human error as patient volumes grow."
p.font.name = "Calibri"
p.font.size = Pt(11.5)
p.font.color.rgb = SLATE_TEXT

# -------------------------------------------------------------
# 2. RIGHT CARD: OBJECTIVE
# -------------------------------------------------------------
card2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.60), Inches(1.70), Inches(6.23), Inches(2.20))
card2.fill.solid()
card2.fill.fore_color.rgb = WHITE
card2.line.color.rgb = BORDER_GRAY
card2.line.width = Pt(1.5)

# Circle Icon for Objective (self-contained with centered text)
icon2 = s2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.95), Inches(1.92), Inches(0.48), Inches(0.48))
icon2.fill.solid()
icon2.fill.fore_color.rgb = TEAL_PRIMARY
icon2.line.fill.background()
tf2 = icon2.text_frame
tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf2.paragraphs[0]
p.text = "→"
p.font.name = "Calibri"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER

# Title for Objective
t2 = s2.shapes.add_textbox(Inches(7.58), Inches(1.92), Inches(5.00), Inches(0.48))
tf_t2 = t2.text_frame
tf_t2.vertical_anchor = MSO_ANCHOR.MIDDLE
tf_t2.margin_left = Inches(0.05)
p = tf_t2.paragraphs[0]
p.text = "Objective"
p.font.name = "Cambria"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = DARK_NAVY

# Body Text for Objective
b2 = s2.shapes.add_textbox(Inches(6.95), Inches(2.55), Inches(5.60), Inches(1.20))
tf_b2 = b2.text_frame
tf_b2.word_wrap = True
tf_b2.margin_left = Inches(0.0)
tf_b2.margin_right = Inches(0.0)
tf_b2.margin_top = Inches(0.0)
p = tf_b2.paragraphs[0]
p.text = "Develop an automated NLP and Ensemble Machine Learning system that predicts the correct medical specialty from a clinical transcription, reducing administrative overhead and speeding up document routing."
p.font.name = "Calibri"
p.font.size = Pt(11.5)
p.font.color.rgb = SLATE_TEXT

# Close PowerPoint if running
subprocess.run(["powershell", "-Command", "Stop-Process -Name POWERPNT -Force -ErrorAction SilentlyContinue"], capture_output=True)
time.sleep(1)

# Save
prs.save(PPTX_PATH)
print(f"Saved workspace PPTX: {PPTX_PATH}")

try:
    prs.save(r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx")
    print("Saved Downloads PPTX successfully.")
except Exception as e:
    print(f"Downloads note: {e}")

print("Slide 2 card positioning fixed cleanly!")
