import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
prs = Presentation(r'D:\Antigravity Projects\Mini Project S3 MCA\model training project\2nd Presentation 30-9-2026\Medical_Specialty_Classification_Interim_Presentation_2.pptx')
for idx in [3, 4]:
    s = prs.slides[idx]
    print(f'=== SLIDE {idx+1} ===')
    for shp in s.shapes:
        left = shp.left / 914400 if shp.left is not None else 0
        top = shp.top / 914400 if shp.top is not None else 0
        w = shp.width / 914400 if shp.width is not None else 0
        h = shp.height / 914400 if shp.height is not None else 0
        txt = shp.text_frame.text[:35].replace('\n', ' ') if shp.has_text_frame else ''
        print(f'  Shape {shp.name}: left={left:.2f}, top={top:.2f}, w={w:.2f}, h={h:.2f} | text="{txt}"')
