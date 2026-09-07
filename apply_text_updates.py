import os
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

pptx_path = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation.pptx"
downloads_path = r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx"

prs = pptx.Presentation(pptx_path)

PURE_BLACK = RGBColor(0, 0, 0)
GREY_HEX_TARGETS = {
    '334155', '5B6472', '64748B', '475569', '718096', '6B7280', 
    '94A3B8', '4B5563', '1E293B', '16213E', '555555', '666666', 
    '4A5568', '707070', '888888', '404040', '2D3748', '374151',
    '9CA3AF'
}

def is_grey_color(col):
    if not col or col.type != 1:
        return False
    hex_code = str(col.rgb).upper()
    return hex_code in GREY_HEX_TARGETS or hex_code.startswith('33') or hex_code.startswith('5B') or hex_code.startswith('64')

# ==========================================
# 1. SPECIAL TREATMENT FOR SLIDE 3 (INDEX 2)
# ==========================================
s3 = prs.slides[2]

# Shapes 2, 3, 4, 5: Top 4 Pill Boxes
pill_accents = {
    2: RGBColor(14, 124, 134),   # Teal
    3: RGBColor(37, 99, 235),    # Blue
    4: RGBColor(147, 51, 234),   # Purple
    5: RGBColor(22, 163, 74)     # Green
}

for sh_idx in [2, 3, 4, 5]:
    sh = s3.shapes[sh_idx]
    if sh.has_text_frame:
        tf = sh.text_frame
        # P0: Phase Title
        if len(tf.paragraphs) > 0:
            p0 = tf.paragraphs[0]
            p0.font.size = Pt(11.5)
            p0.font.bold = True
            p0.font.color.rgb = pill_accents[sh_idx]
            for r in p0.runs:
                r.font.size = Pt(11.5)
                r.font.bold = True
                r.font.color.rgb = pill_accents[sh_idx]
        # P1: Phase Subtitle
        if len(tf.paragraphs) > 1:
            p1 = tf.paragraphs[1]
            p1.font.size = Pt(11.0)
            p1.font.bold = True
            p1.font.color.rgb = PURE_BLACK
            for r in p1.runs:
                r.font.size = Pt(11.0)
                r.font.bold = True
                r.font.color.rgb = PURE_BLACK

# Shapes 6, 7, 8, 9: The 4 Big Cards
card_accents = {
    6: RGBColor(14, 124, 134),   # Clinical NLP
    7: RGBColor(37, 99, 235),    # TF-IDF
    8: RGBColor(22, 163, 74),    # Supervised ML
    9: RGBColor(147, 51, 234)    # Ensemble & Gap
}

for sh_idx in [6, 7, 8, 9]:
    sh = s3.shapes[sh_idx]
    if sh.has_text_frame:
        tf = sh.text_frame
        tf.margin_top = Inches(0.16)
        tf.margin_bottom = Inches(0.12)
        tf.margin_left = Inches(0.22)
        tf.margin_right = Inches(0.22)
        
        # P0: Card Title
        if len(tf.paragraphs) > 0:
            p0 = tf.paragraphs[0]
            p0.font.size = Pt(15.0)
            p0.font.bold = True
            p0.font.color.rgb = card_accents[sh_idx]
            p0.space_after = Pt(7)
            for r in p0.runs:
                r.font.size = Pt(15.0)
                r.font.bold = True
                r.font.color.rgb = card_accents[sh_idx]
        
        # P1, P2, P3: Bullets
        for p_idx in range(1, len(tf.paragraphs)):
            p = tf.paragraphs[p_idx]
            p.space_after = Pt(5.5)
            # Run 0 is bullet symbol, Run 1 is text
            for r_idx, r in enumerate(p.runs):
                if r_idx == 0:
                    r.font.size = Pt(13.0)
                    r.font.bold = True
                    r.font.color.rgb = card_accents[sh_idx]
                else:
                    r.font.size = Pt(12.5)
                    r.font.bold = False
                    r.font.color.rgb = PURE_BLACK

print("Slide 3 updated successfully!")

# ========================================================
# 2. GENERAL PASS ACROSS ALL LIGHT SLIDES (SLIDES 2 TO 18)
# ========================================================
# We skip Slide 1 (index 0) and Slide 19 (index 18) because they are dark theme.
for s_idx in range(1, 18):
    s = prs.slides[s_idx]
    if s_idx == 2:
        # Already customized Slide 3, but update footers
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.color and is_grey_color(r.font.color):
                            r.font.color.rgb = PURE_BLACK
                            if r.font.size and r.font.size.pt < 10.0:
                                r.font.size = Pt(10.0)
        continue
        
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        for p in tf.paragraphs:
            # Check paragraph font
            if p.font.color and is_grey_color(p.font.color):
                p.font.color.rgb = PURE_BLACK
                if p.font.size:
                    if p.font.size.pt <= 9.5:
                        p.font.size = Pt(min(p.font.size.pt + 1.5, 11.0))
                    elif p.font.size.pt < 12.0:
                        p.font.size = Pt(min(p.font.size.pt + 1.5, 12.5))
            
            # Check runs
            for r in p.runs:
                if r.font.color and is_grey_color(r.font.color):
                    r.font.color.rgb = PURE_BLACK
                    if r.font.size:
                        if r.font.size.pt <= 9.5:
                            r.font.size = Pt(min(r.font.size.pt + 1.5, 11.0))
                        elif r.font.size.pt < 12.0:
                            r.font.size = Pt(min(r.font.size.pt + 1.5, 12.5))
                # Also if run has size < 10.5 and font color is not set (inheriting dark), bump size
                elif (not r.font.color or r.font.color.type != 1) and r.font.size and r.font.size.pt < 10.5:
                    r.font.size = Pt(min(r.font.size.pt + 1.5, 11.5))

# ========================================================
# 3. SPECIFIC HIGH-VISIBILITY ADJUSTMENTS FOR KEY SLIDES
# ========================================================

# Slide 2 (Problem & Objectives)
s2 = prs.slides[1]
# Shape 30: Problem Statement Body
if len(s2.shapes) > 30 and s2.shapes[30].has_text_frame:
    p = s2.shapes[30].text_frame.paragraphs[0]
    p.font.size = Pt(13.0)
    p.font.color.rgb = PURE_BLACK
    for r in p.runs:
        r.font.size = Pt(13.0)
        r.font.color.rgb = PURE_BLACK

# Shape 34: Objective Statement Body
if len(s2.shapes) > 34 and s2.shapes[34].has_text_frame:
    p = s2.shapes[34].text_frame.paragraphs[0]
    p.font.size = Pt(13.0)
    p.font.color.rgb = PURE_BLACK
    for r in p.runs:
        r.font.size = Pt(13.0)
        r.font.color.rgb = PURE_BLACK

# Shapes 11, 16, 21, 26: 4 Objective item descriptions
for sh_idx in [11, 16, 21, 26]:
    if len(s2.shapes) > sh_idx and s2.shapes[sh_idx].has_text_frame:
        p = s2.shapes[sh_idx].text_frame.paragraphs[0]
        p.font.size = Pt(11.5)
        p.font.color.rgb = PURE_BLACK
        for r in p.runs:
            r.font.size = Pt(11.5)
            r.font.color.rgb = PURE_BLACK

# Shapes 10, 15, 20, 25: 4 Objective item titles
for sh_idx in [10, 15, 20, 25]:
    if len(s2.shapes) > sh_idx and s2.shapes[sh_idx].has_text_frame:
        p = s2.shapes[sh_idx].text_frame.paragraphs[0]
        p.font.size = Pt(13.5)
        p.font.bold = True

# Slide 4 (Summary of Key Research Papers)
s4 = prs.slides[3]
for sh_idx in [11, 13, 15, 21, 23, 25, 31, 33, 35]:
    if len(s4.shapes) > sh_idx and s4.shapes[sh_idx].has_text_frame:
        p = s4.shapes[sh_idx].text_frame.paragraphs[0]
        p.font.size = Pt(12.0)
        p.font.color.rgb = PURE_BLACK
        for r in p.runs:
            r.font.size = Pt(12.0)
            r.font.color.rgb = PURE_BLACK

prs.save(pptx_path)
print("Saved presentation to:", pptx_path)

# Copy to downloads
shutil.copy2(pptx_path, downloads_path)
print("Synced presentation to:", downloads_path)
