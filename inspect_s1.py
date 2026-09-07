import pptx

prs = pptx.Presentation(r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx")
s1 = prs.slides[0]
for idx, sh in enumerate(s1.shapes):
    t = sh.text.strip().replace('\n', ' ') if sh.has_text_frame else ""
    print(f"{idx}: {sh.name} | left={sh.left/914400:.2f} in, top={sh.top/914400:.2f} in, w={sh.width/914400:.2f} in, h={sh.height/914400:.2f} in | text={t[:40]}")
