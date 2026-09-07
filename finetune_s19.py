import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PPTX_PATH = r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation_updated.pptx"
prs = pptx.Presentation(PPTX_PATH)

TEAL_ACCENT = RGBColor(0x0E, 0x7C, 0x86)
CYAN_BRIGHT = RGBColor(0x00, 0xE5, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_NAVY = RGBColor(0x0B, 0x1F, 0x3A)

# Slide 19 check
s19 = prs.slides[18]
# Check if subbanner exists
has_subbanner = False
for sh in s19.shapes:
    if sh.has_text_frame and "ACADEMIC TIMELINE" in sh.text:
        has_subbanner = True
        break

if not has_subbanner:
    sb = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.50), Inches(1.50), Inches(12.30), Inches(0.40))
    sb.fill.solid()
    sb.fill.fore_color.rgb = RGBColor(0xEE, 0xF6, 0xF6)
    sb.line.color.rgb = TEAL_ACCENT
    sb.line.width = Pt(1)
    tf = sb.text_frame
    tf.margin_top = Inches(0.06)
    tf.margin_bottom = Inches(0.06)
    p = tf.paragraphs[0]
    p.text = "ACADEMIC TIMELINE & PRESENTATION MILESTONES (SEMESTER 3 MCA MINI PROJECT)"
    p.font.name = "Calibri"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEAL_ACCENT
    p.alignment = PP_ALIGN.CENTER
    print("Slide 19 subbanner added.")

prs.save(PPTX_PATH)
prs.save(r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_updated.pptx")
print("Verified and saved updated PPTX!")
