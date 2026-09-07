import pptx
from pptx.dml.color import RGBColor
from pptx.util import Pt

SRC_PATH = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_v4.pptx"
prs = pptx.Presentation(SRC_PATH)

slide_idx_to_remove = None
for idx, s in enumerate(prs.slides):
    for sh in s.shapes:
        if sh.has_text_frame and "Clinical Text Data Cleaning Pipeline" in sh.text:
            slide_idx_to_remove = idx
            break
    if slide_idx_to_remove is not None:
        break

print(f"Target slide to remove: index {slide_idx_to_remove} (Slide {slide_idx_to_remove + 1})")

if slide_idx_to_remove is not None:
    # Remove slide from presentation
    rId = prs.slides._sldIdLst[slide_idx_to_remove].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[slide_idx_to_remove]
    print("Slide removed successfully.")

total_slides = len(prs.slides)
print(f"Total slides remaining: {total_slides}")

# Update footers across all slides
for idx, slide in enumerate(prs.slides):
    slide_num = idx + 1
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = sh.text.strip()
            if "Slide " in t and " of " in t:
                sh.text = f"Slide {slide_num} of {total_slides}"
                p = sh.text_frame.paragraphs[0]
                p.font.name = "Calibri"
                p.font.size = Pt(9)
                p.font.color.rgb = RGBColor(0x5B, 0x64, 0x72) if idx != 0 and idx != total_slides - 1 else RGBColor(0x94, 0xA3, 0xB8)

# Save to v5
V5_LOCAL = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_v5.pptx"
V5_DOWNLOADS = r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation_v5.pptx"

prs.save(V5_LOCAL)
print(f"Saved: {V5_LOCAL}")
prs.save(V5_DOWNLOADS)
print(f"Saved: {V5_DOWNLOADS}")

# Also update previous filenames if unlocked
for path in [
    r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx",
    r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation_updated.pptx",
    r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_updated.pptx",
    r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_v4.pptx"
]:
    try:
        prs.save(path)
        print(f"Also updated: {path}")
    except Exception as e:
        print(f"Note: {path} is locked by PowerPoint: {e}")

print("Done!")
