import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

SRC_PATH = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_v3.pptx"
prs = pptx.Presentation(SRC_PATH)

# Colors
TEAL_PRIMARY = RGBColor(0x0E, 0x7C, 0x86)
DARK_NAVY = RGBColor(0x0B, 0x1F, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE_DARK = RGBColor(0x33, 0x41, 0x55)
BORDER_GRAY = RGBColor(0xCB, 0xD5, 0xE1)
CYAN_ACCENT = RGBColor(0x00, 0xE5, 0xFF)
CYAN_LIGHT = RGBColor(0x38, 0xBD, 0xF8)
PURPLE_ACCENT = RGBColor(0x93, 0x33, 0xEA)
GREEN_ACCENT = RGBColor(0x16, 0xA3, 0x4A)
BLUE_ACCENT = RGBColor(0x25, 0x63, 0xEB)
LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)

# 1. Add new slide for "Literature Review: Background & Research Landscape"
layout = prs.slide_layouts[0]
new_slide = prs.slides.add_slide(layout)

# Clear default shapes on new slide
sp_tree = new_slide.shapes._spTree
for sh in list(new_slide.shapes):
    sp_tree.remove(sh._element)

# Move new slide to index 2 (Slide 3)
sldIdLst = prs.slides._sldIdLst
sldIdLst.insert(2, sldIdLst[-1])

# Build Slide 3: Background & Research Landscape
# Header Pill
pill = new_slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.50), Inches(0.35), Inches(4.60), Inches(0.32))
pill.fill.solid()
pill.fill.fore_color.rgb = TEAL_PRIMARY
pill.line.fill.background()
p = pill.text_frame.paragraphs[0]
p.text = "MCA MINI PROJECT | LITERATURE REVIEW"
p.font.name = "Calibri"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = WHITE

# Slide Title
tbox = new_slide.shapes.add_textbox(Inches(0.50), Inches(0.72), Inches(12.30), Inches(0.68))
p = tbox.text_frame.paragraphs[0]
p.text = "Literature Review: Background & Research Landscape"
p.font.name = "Cambria"
p.font.size = Pt(26)
p.font.bold = True
p.font.color.rgb = DARK_NAVY

# 4 Top Phase Badges
phase_badges = [
    ("1. Prior Literature", "Validation of ML for clinical text", TEAL_PRIMARY, RGBColor(0xF0, 0xFD, 0xFA)),
    ("2. Key Methodologies", "TF-IDF vectorization & linear classifiers", BLUE_ACCENT, RGBColor(0xEF, 0xF6, 0xFF)),
    ("3. Identified Gap", "High compute requirements & model variance", PURPLE_ACCENT, RGBColor(0xFA, 0xF5, 0xFF)),
    ("4. Proposed Architecture", "Lightweight soft voting ensemble on web", GREEN_ACCENT, RGBColor(0xF0, 0xFD, 0xF4))
]

b_w = Inches(2.90)
b_gap = Inches(0.23)
b_start = Inches(0.50)
for i, (p_title, p_sub, col, bg) in enumerate(phase_badges):
    bx = new_slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, b_start + i * (b_w + b_gap), Inches(1.50), b_w, Inches(0.75))
    bx.fill.solid()
    bx.fill.fore_color.rgb = bg
    bx.line.color.rgb = col
    bx.line.width = Pt(1.2)
    tf = bx.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = p_title
    p.font.name = "Calibri"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = col
    p.space_after = Pt(2)
    p = tf.add_paragraph()
    p.text = p_sub
    p.font.name = "Calibri"
    p.font.size = Pt(9)
    p.font.color.rgb = SLATE_DARK

# 4 Main Literature Topic Cards (2x2 Grid)
grid_w = Inches(5.95)
grid_h = Inches(2.20)
x_left = Inches(0.50)
x_right = Inches(6.85)
y_top = Inches(2.40)
y_bot = Inches(4.75)

cards_data = [
    (x_left, y_top, "Clinical NLP & Preprocessing", [
        "Clinical narratives contain non-standard doctor shorthand, abbreviations, and punctuation noise.",
        "Implements a 5-step sequence: Cleaning, Lowercasing, Tokenization, Stop-word removal, and Lemmatization.",
        "Morphological lemmatization normalizes inflected clinical terms while preserving diagnostic root semantics."
    ], TEAL_PRIMARY, RGBColor(0xF0, 0xFD, 0xFA)),
    
    (x_right, y_top, "TF-IDF Feature Representation", [
        "Converts unstructured medical text into standardized, high-dimensional sparse numerical feature vectors.",
        "Sublinear scaling (1 + log(TF)) dampens repetitive non-informative terms while highlighting discriminative tokens.",
        "Computationally lightweight; enables rapid CPU-based real-time inference on web platforms."
    ], BLUE_ACCENT, RGBColor(0xEF, 0xF6, 0xFF)),
    
    (x_left, y_bot, "Supervised ML Classifiers", [
        "Linear models (SVM, Logistic Regression) excel on high-dimensional sparse text representations.",
        "Random Forest introduces non-linear decision partitioning and bagging variance reduction.",
        "Multinomial Naive Bayes provides fast, effective probabilistic likelihood estimation for text vectors."
    ], GREEN_ACCENT, RGBColor(0xF0, 0xFD, 0xF4)),
    
    (x_right, y_bot, "Ensemble Learning & Research Gap", [
        "Identified Gap: Many deep learning approaches demand heavy GPU resources; single models show variance on rare classes.",
        "Proposed Solution: An accessible Soft Voting Ensemble combining calibrated base models for balanced classification.",
        "Practical Outcome: Achieves robust multi-class accuracy across all 40 specialties on standard CPU hardware."
    ], PURPLE_ACCENT, RGBColor(0xFA, 0xF5, 0xFF))
]

for x, y, title, bullets, col, bg in cards_data:
    card = new_slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, grid_w, grid_h)
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = col
    card.line.width = Pt(1.5)
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.20)
    tf.margin_right = Inches(0.20)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.10)
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = col
    p.space_after = Pt(4)
    
    for b in bullets:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = "• "
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = col
        
        r2 = p.add_run()
        r2.text = b
        r2.font.name = "Calibri"
        r2.font.size = Pt(9)
        r2.font.color.rgb = SLATE_DARK
        p.space_after = Pt(2)

print("Slide 3 (Literature Review Background) created.")

# 2. Update Slide 4 (formerly Slide 3): Title should be "Summary of Key Research Papers"
s4 = prs.slides[3]
for sh in s4.shapes:
    if sh.has_text_frame:
        t = sh.text.strip()
        if t == "Literature Review":
            sh.text = "Summary of Key Research Papers"
            p = sh.text_frame.paragraphs[0]
            p.font.name = "Cambria"
            p.font.size = Pt(26)
            p.font.bold = True
            p.font.color.rgb = DARK_NAVY
        elif "SUPPORTING LITERATURE" in t:
            sh.text = "MCA MINI PROJECT | LITERATURE REVIEW"
            p = sh.text_frame.paragraphs[0]
            p.font.name = "Calibri"
            p.font.size = Pt(10.5)
            p.font.bold = True
            p.font.color.rgb = WHITE

print("Slide 4 title updated to 'Summary of Key Research Papers'.")

# 3. Update all slide footers to "Slide X of 21"
total_slides = len(prs.slides)
print(f"Total slides now: {total_slides}")

for idx, slide in enumerate(prs.slides):
    slide_num = idx + 1
    # Check footers on slide
    footer_found = False
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = sh.text.strip()
            if "Slide " in t and " of " in t:
                sh.text = f"Slide {slide_num} of {total_slides}"
                p = sh.text_frame.paragraphs[0]
                p.font.name = "Calibri"
                p.font.size = Pt(9)
                p.font.color.rgb = RGBColor(0x5B, 0x64, 0x72) if idx != 0 and idx != total_slides - 1 else RGBColor(0x94, 0xA3, 0xB8)
                footer_found = True
    
    # If this is our newly inserted slide 3, add footers
    if idx == 2 and not footer_found:
        f1 = slide.shapes.add_textbox(Inches(0.50), Inches(7.13), Inches(4.60), Inches(0.30))
        p1 = f1.text_frame.paragraphs[0]
        p1.text = "Department of Computer Applications | MACE"
        p1.font.name = "Calibri"
        p1.font.size = Pt(9)
        p1.font.color.rgb = RGBColor(0x5B, 0x64, 0x72)

        f2 = slide.shapes.add_textbox(Inches(4.00), Inches(7.13), Inches(5.33), Inches(0.30))
        p2 = f2.text_frame.paragraphs[0]
        p2.text = "Mar Athanasius College of Engineering, Kothamangalam"
        p2.font.name = "Calibri"
        p2.font.size = Pt(9)
        p2.font.color.rgb = RGBColor(0x5B, 0x64, 0x72)

        f3 = slide.shapes.add_textbox(Inches(9.90), Inches(7.13), Inches(2.93), Inches(0.30))
        p3 = f3.text_frame.paragraphs[0]
        p3.text = f"Slide 3 of {total_slides}"
        p3.font.name = "Calibri"
        p3.font.size = Pt(9)
        p3.font.color.rgb = RGBColor(0x5B, 0x64, 0x72)

# Save to all destination paths
V4_LOCAL = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_v4.pptx"
V4_DOWNLOADS = r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation_v4.pptx"

prs.save(V4_LOCAL)
print(f"Saved: {V4_LOCAL}")
prs.save(V4_DOWNLOADS)
print(f"Saved: {V4_DOWNLOADS}")

# Try updating the previous paths if unlocked
for p in [r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation_updated.pptx", r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx", r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_updated.pptx"]:
    try:
        prs.save(p)
        print(f"Also updated: {p}")
    except Exception as e:
        print(f"Note: {p} locked by PowerPoint.")

print("All tasks completed successfully.")
