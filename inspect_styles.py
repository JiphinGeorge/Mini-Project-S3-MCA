import pptx

prs = pptx.Presentation(r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx")

for s_idx in [1, 12, 16]:  # Slide 2, 13, 17
    s = prs.slides[s_idx]
    print(f"=== SLIDE {s_idx+1} ===")
    for sh in s.shapes:
        line_color = None
        fill_color = None
        try:
            if sh.line and sh.line.color and sh.line.color.type == 1:
                line_color = sh.line.color.rgb
        except: pass
        try:
            if sh.fill and sh.fill.type == 1 and sh.fill.fore_color:
                fill_color = sh.fill.fore_color.rgb
        except: pass
        
        font_info = ""
        if sh.has_text_frame and sh.text_frame.paragraphs:
            p = sh.text_frame.paragraphs[0]
            if p.runs:
                r = p.runs[0]
                fc = r.font.color.rgb if r.font.color and r.font.color.type == 1 else None
                font_info = f"font={r.font.name} size={r.font.size} color={fc} bold={r.font.bold}"
        print(f"  {sh.name} ({sh.shape_type}): left={sh.left} top={sh.top} w={sh.width} h={sh.height} fill={fill_color} line={line_color} | {font_info} | text={repr(sh.text[:30] if sh.has_text_frame else '')}")
