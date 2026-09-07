import pptx

SRC_PATH = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_v7.pptx"
prs = pptx.Presentation(SRC_PATH)

s12 = prs.slides[11]
sp_tree = s12.shapes._spTree

for sh in list(s12.shapes):
    if sh.top >= pptx.util.Inches(5.6) and sh.top < pptx.util.Inches(6.9) and sh.width > pptx.util.Inches(8.0):
        print(f"Removing residual box: {sh.name} at top={sh.top/914400:.2f} in")
        sp_tree.remove(sh._element)

# Now check all shapes
print("Remaining shapes above footer on Slide 12:")
for sh in s12.shapes:
    if sh.top > pptx.util.Inches(4.5):
        t = sh.text.strip().replace('\n', ' ') if sh.has_text_frame else ''
        print(f"  {sh.name}: top={sh.top/914400:.2f} in, h={sh.height/914400:.2f} in | text={repr(t[:35])}")

# Save
for p in [
    SRC_PATH,
    r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation_v7.pptx",
    r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx",
    r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation_updated.pptx"
]:
    try:
        prs.save(p)
        print(f"Saved: {p}")
    except Exception as e:
        print(f"Locked: {p}")

print("Cleaned up successfully!")
