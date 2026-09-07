import pptx

prs = pptx.Presentation(r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx")
print(f"Total slides: {len(prs.slides)}")

for idx in [0, 12, 17, 18, 19]:
    s = prs.slides[idx]
    print(f"\n=================== SLIDE {idx+1} ===================")
    for sh in s.shapes:
        if sh.has_text_frame and sh.text.strip():
            print(f"[{sh.name}] left={sh.left/914400:.2f} in, top={sh.top/914400:.2f} in, w={sh.width/914400:.2f} in, h={sh.height/914400:.2f} in:")
            for p in sh.text_frame.paragraphs[:4]:
                if p.text.strip():
                    print(f"   {p.text.strip()[:65]}")
