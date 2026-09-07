import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PPTX_PATH = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation.pptx"
prs = pptx.Presentation(PPTX_PATH)

s12 = prs.slides[11]
sp_tree = s12.shapes._spTree

# Remove the broken bottom shapes (Shape 26 to 34)
# Let's collect any shapes with top > 4.3 in and top < 7.0 in
shapes_to_remove = []
for sh in s12.shapes:
    if sh.top > Inches(4.3) and sh.top < Inches(7.0):
        shapes_to_remove.append(sh)

print(f"Removing {len(shapes_to_remove)} misaligned shapes...")
for sh in shapes_to_remove:
    sp_tree.remove(sh._element)

# Now add the 3 perfectly styled, unified cards:
card_y = Inches(4.75)
card_h = Inches(1.75)

# 1. Feature Extraction Card
c1_x = Inches(0.50)
c1_w = Inches(2.70)
card1 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_x, card_y, c1_w, card_h)
card1.fill.solid()
card1.fill.fore_color.rgb = RGBColor(0x13, 0x29, 0x4B) # Deep Navy
card1.line.color.rgb = RGBColor(0x00, 0xE5, 0xFF)      # Cyan border
card1.line.width = Pt(1.5)

tf1 = card1.text_frame
tf1.word_wrap = True
tf1.margin_left = Inches(0.18)
tf1.margin_right = Inches(0.18)
tf1.margin_top = Inches(0.18)
tf1.margin_bottom = Inches(0.15)

p = tf1.paragraphs[0]
p.text = "FEATURE EXTRACTION"
p.font.name = "Calibri"
p.font.size = Pt(10)
p.font.bold = True
p.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8) # Cyan label
p.space_after = Pt(4)

p = tf1.add_paragraph()
p.text = "TF-IDF Vectorizer"
p.font.name = "Cambria"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) # White
p.space_after = Pt(4)

p = tf1.add_paragraph()
p.text = "Unigrams & bigrams with sublinear TF scaling and L2 normalization."
p.font.name = "Calibri"
p.font.size = Pt(9.5)
p.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

# 2. Classification Card
c2_x = Inches(3.40)
c2_w = Inches(6.15)
card2 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_x, card_y, c2_w, card_h)
card2.fill.solid()
card2.fill.fore_color.rgb = RGBColor(0x0E, 0x7C, 0x86) # Teal
card2.line.color.rgb = RGBColor(0x2D, 0xD4, 0xBF)      # Emerald border
card2.line.width = Pt(1.5)

tf2 = card2.text_frame
tf2.word_wrap = True
tf2.margin_left = Inches(0.22)
tf2.margin_right = Inches(0.22)
tf2.margin_top = Inches(0.18)
tf2.margin_bottom = Inches(0.15)

p = tf2.paragraphs[0]
p.text = "CLASSIFICATION (4 DIVERSE BASE ESTIMATORS)"
p.font.name = "Calibri"
p.font.size = Pt(10)
p.font.bold = True
p.font.color.rgb = RGBColor(0xA7, 0xF3, 0xD0) # Light mint label
p.space_after = Pt(4)

p = tf2.add_paragraph()
p.text = "SVM + Random Forest + Logistic Regression + Multinomial NB"
p.font.name = "Cambria"
p.font.size = Pt(15.5)
p.font.bold = True
p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) # White
p.space_after = Pt(4)

p = tf2.add_paragraph()
p.text = "Each model independently evaluates high-dimensional sparse vectors and outputs class probabilities."
p.font.name = "Calibri"
p.font.size = Pt(9.5)
p.font.color.rgb = RGBColor(0xF0, 0xFD, 0xFA)

# 3. Ensemble Card
c3_x = Inches(9.75)
c3_w = Inches(3.08)
card3 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c3_x, card_y, c3_w, card_h)
card3.fill.solid()
card3.fill.fore_color.rgb = RGBColor(0x5B, 0x21, 0xB6) # Rich Purple
card3.line.color.rgb = RGBColor(0xC0, 0x84, 0xFC)      # Lavender border
card3.line.width = Pt(1.5)

tf3 = card3.text_frame
tf3.word_wrap = True
tf3.margin_left = Inches(0.18)
tf3.margin_right = Inches(0.18)
tf3.margin_top = Inches(0.18)
tf3.margin_bottom = Inches(0.15)

p = tf3.paragraphs[0]
p.text = "ENSEMBLE AGGREGATION"
p.font.name = "Calibri"
p.font.size = Pt(10)
p.font.bold = True
p.font.color.rgb = RGBColor(0xE9, 0xD5, 0xFF) # Lavender label
p.space_after = Pt(4)

p = tf3.add_paragraph()
p.text = "Voting Classifier (Soft)"
p.font.name = "Cambria"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) # White
p.space_after = Pt(4)

p = tf3.add_paragraph()
p.text = "Aggregates consensus probabilities to output final specialty prediction."
p.font.name = "Calibri"
p.font.size = Pt(9.5)
p.font.color.rgb = RGBColor(0xF5, 0xF3, 0xFF)

# Close PowerPoint if running to release file lock
import subprocess, time
subprocess.run(["powershell", "-Command", "Stop-Process -Name POWERPNT -Force -ErrorAction SilentlyContinue"], capture_output=True)
time.sleep(1)

# Save to single workspace file
prs.save(PPTX_PATH)
print(f"Successfully saved to: {PPTX_PATH}")

# Also update Downloads
try:
    prs.save(r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx")
    print("Also updated Downloads file.")
except Exception as e:
    print(f"Downloads file note: {e}")

print("Card fix complete!")
