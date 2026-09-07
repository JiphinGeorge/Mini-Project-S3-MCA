import pptx
from pptx.util import Inches, Pt
import shutil
import os

SRC_PATH = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_updated.pptx"
prs = pptx.Presentation(SRC_PATH)

s1 = prs.slides[0]
sp_tree = s1.shapes._spTree

# Remove Shape 4 (TextBox 5 with subtitle "Clinical NLP & Supervised Multi-Class Learning on MTSamples Medical Corpus")
shape_to_remove = None
for sh in s1.shapes:
    if sh.has_text_frame and "Clinical NLP & Supervised Multi-Class Learning" in sh.text:
        shape_to_remove = sh
        break

if shape_to_remove:
    sp_tree.remove(shape_to_remove._element)
    print("Subtitle removed successfully.")

# Adjust cards and badges
for sh in s1.shapes:
    if sh.has_text_frame:
        t = sh.text.strip()
        if "PRESENTED BY" in t:
            sh.top = Inches(3.65)
            sh.height = Inches(2.45)
            # Polish paragraph spacing
            for p in sh.text_frame.paragraphs:
                if "Jiphin George" in p.text:
                    p.font.size = Pt(24)
                    p.space_after = Pt(6)
                elif "PRESENTED BY" in p.text:
                    p.space_after = Pt(6)
                elif "MAC25MCA-2033" in p.text:
                    p.space_after = Pt(4)
        elif "PROJECT GUIDE" in t:
            sh.top = Inches(3.65)
            sh.height = Inches(2.45)
            for p in sh.text_frame.paragraphs:
                if "Prof. Biju Skaria" in p.text:
                    p.font.size = Pt(24)
                    p.space_after = Pt(6)
                elif "PROJECT GUIDE" in p.text:
                    p.space_after = Pt(6)
                elif "Department of Computer" in p.text:
                    p.space_after = Pt(4)
        # Check badges
        elif t in ["Clinical NLP", "TF-IDF Vectorization", "Supervised ML", "Soft Voting Ensemble", "Flask Deployment"]:
            sh.top = Inches(6.35)
            sh.height = Inches(0.45)

# Save to v3 files
V3_LOCAL = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_v3.pptx"
V3_DOWNLOADS = r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation_v3.pptx"

prs.save(V3_LOCAL)
print(f"Saved successfully to: {V3_LOCAL}")

prs.save(V3_DOWNLOADS)
print(f"Saved successfully to: {V3_DOWNLOADS}")

# Also try to save back to updated and original if unlocked
for p in [SRC_PATH, r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation_updated.pptx", r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx"]:
    try:
        prs.save(p)
        print(f"Also updated: {p}")
    except Exception as e:
        print(f"Note: {p} is currently open in PowerPoint.")

print("All done!")

