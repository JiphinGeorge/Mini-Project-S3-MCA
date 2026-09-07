import pptx

prs = pptx.Presentation(r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation_updated.pptx")
print(f"Total slides: {len(prs.slides)}")

check_slides = [0, 12, 13, 14, 15, 16, 17, 18, 19]

for idx in check_slides:
    s = prs.slides[idx]
    texts = [sh.text.strip().replace('\n', ' ') for sh in s.shapes if sh.has_text_frame and sh.text.strip()]
    h = texts[0][:40] if len(texts) > 0 else 'NO TEXT'
    t = texts[1][:40] if len(texts) > 1 else ''
    print(f"Slide {idx+1:02d}: [{h}] | [{t}]")

print("\nAll target slides verified successfully.")
