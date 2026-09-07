import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Let's inspect the slides
prs = pptx.Presentation(r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_v3.pptx")
print(f"Current slides: {len(prs.slides)}")
for i, s in enumerate(prs.slides):
    t = ""
    for sh in s.shapes:
        if sh.has_text_frame and sh.text.strip():
            t = sh.text.strip().replace('\n', ' ')[:40]
            break
    print(f"Slide {i+1}: {t}")
