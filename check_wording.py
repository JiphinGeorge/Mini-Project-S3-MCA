import pptx

prs = pptx.Presentation(r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx")
with open("target_svm_check.txt", "w", encoding="utf-8") as f:
    for idx, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame:
                t = sh.text
                if "Linear SVM" in t or "functional margin" in t or "Assistant Professor" in t:
                    f.write(f"Slide {idx+1}: [{sh.name}] -> {t}\n")
print("Check completed.")
