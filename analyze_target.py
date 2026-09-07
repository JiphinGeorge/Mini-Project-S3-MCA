import pptx

prs = pptx.Presentation(r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx")
with open("target_full_analysis.txt", "w", encoding="utf-8") as f:
    for i, slide in enumerate(prs.slides):
        f.write(f"\n==================== SLIDE {i+1} ====================\n")
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text.strip():
                f.write(f"[{sh.name}] (left={sh.left/914400:.2f}, top={sh.top/914400:.2f}, w={sh.width/914400:.2f}, h={sh.height/914400:.2f}):\n")
                f.write(f"  {repr(sh.text.strip())}\n")
print("Target analysis written successfully.")
