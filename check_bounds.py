import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
prs = Presentation(r'D:\Antigravity Projects\Mini Project S3 MCA\model training project\2nd Presentation 30-9-2026\Medical_Specialty_Classification_Interim_Presentation_2.pptx')
for idx, s in enumerate(prs.slides):
    errors = []
    for shp in s.shapes:
        left = shp.left / 914400 if shp.left is not None else 0
        top = shp.top / 914400 if shp.top is not None else 0
        w = shp.width / 914400 if shp.width is not None else 0
        h = shp.height / 914400 if shp.height is not None else 0
        
        if left > 13.33 or top > 7.5 or (left + w) > 13.4 or (top + h) > 7.5:
            errors.append(f'OUT OF BOUNDS: {shp.name} at left={left:.2f}, top={top:.2f}, w={w:.2f}, h={h:.2f}')
    if errors:
        print(f'Slide {idx+1} has {len(errors)} errors:')
        for e in errors:
            print('  ', e)
    else:
        print(f'Slide {idx+1}: OK')
