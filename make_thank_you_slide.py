import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

SRC_PATH = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_v7.pptx"
prs = pptx.Presentation(SRC_PATH)

# The final slide is the last slide in prs.slides
s_last = prs.slides[-1]
sp_tree = s_last.shapes._spTree

# Clear all shapes on the final slide to rebuild cleanly
for sh in list(s_last.shapes):
    sp_tree.remove(sh._element)

# Slide Background: Deep Royal Navy
s_last.background.fill.solid()
s_last.background.fill.fore_color.rgb = RGBColor(0x0A, 0x19, 0x2F)

# Top accent bar
bar = s_last.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
bar.fill.solid()
bar.fill.fore_color.rgb = RGBColor(0x00, 0xE5, 0xFF)
bar.line.fill.background()

# Central Hero "THANK YOU" Card
card_w = Inches(10.50)
card_h = Inches(4.80)
card_x = Inches(1.416)
card_y = Inches(1.35)

ty_card = s_last.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y, card_w, card_h)
ty_card.fill.solid()
ty_card.fill.fore_color.rgb = RGBColor(0x11, 0x22, 0x40)
ty_card.line.color.rgb = RGBColor(0x00, 0xE5, 0xFF)
ty_card.line.width = Pt(2)

tf = ty_card.text_frame
tf.word_wrap = True
tf.margin_top = Inches(0.40)
tf.margin_bottom = Inches(0.30)
tf.margin_left = Inches(0.40)
tf.margin_right = Inches(0.40)

# "THANK YOU!" Title
p = tf.paragraphs[0]
p.text = "THANK YOU!"
p.font.name = "Cambria"
p.font.size = Pt(46)
p.font.bold = True
p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
p.alignment = PP_ALIGN.CENTER
p.space_after = Pt(8)

# "Questions & Discussion" Subtitle
p = tf.add_paragraph()
p.text = "Questions, Suggestions & Discussion"
p.font.name = "Calibri"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
p.alignment = PP_ALIGN.CENTER
p.space_after = Pt(18)

# Project Title
p = tf.add_paragraph()
p.text = "Medical Specialty Classification Using TF-IDF\nand Ensemble Machine Learning"
p.font.name = "Cambria"
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = RGBColor(0xF1, 0xF5, 0xF9)
p.alignment = PP_ALIGN.CENTER
p.space_after = Pt(20)

# Presenter & Guide Credentials
p = tf.add_paragraph()
p.text = "Presented by: Jiphin George (Register No: MAC25MCA-2033)"
p.font.name = "Calibri"
p.font.size = Pt(13.5)
p.font.bold = True
p.font.color.rgb = RGBColor(0x4A, 0xDE, 0x80)
p.alignment = PP_ALIGN.CENTER
p.space_after = Pt(4)

p = tf.add_paragraph()
p.text = "Project Guide: Prof. Biju Skaria  |  Department of Computer Applications"
p.font.name = "Calibri"
p.font.size = Pt(12.5)
p.font.bold = False
p.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
p.alignment = PP_ALIGN.CENTER
p.space_after = Pt(4)

p = tf.add_paragraph()
p.text = "Mar Athanasius College of Engineering, Kothamangalam  |  MCA Mini Project (Semester 3)"
p.font.name = "Calibri"
p.font.size = Pt(11.5)
p.font.bold = False
p.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
p.alignment = PP_ALIGN.CENTER

# Footers
total_slides = len(prs.slides)
f1 = s_last.shapes.add_textbox(Inches(0.50), Inches(7.13), Inches(4.60), Inches(0.30))
p1 = f1.text_frame.paragraphs[0]
p1.text = "Department of Computer Applications | MACE"
p1.font.name = "Calibri"
p1.font.size = Pt(9)
p1.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

f2 = s_last.shapes.add_textbox(Inches(4.00), Inches(7.13), Inches(5.33), Inches(0.30))
p2 = f2.text_frame.paragraphs[0]
p2.text = "Mar Athanasius College of Engineering, Kothamangalam"
p2.font.name = "Calibri"
p2.font.size = Pt(9)
p2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

f3 = s_last.shapes.add_textbox(Inches(9.90), Inches(7.13), Inches(2.93), Inches(0.30))
p3 = f3.text_frame.paragraphs[0]
p3.text = f"Slide {total_slides} of {total_slides}"
p3.font.name = "Calibri"
p3.font.size = Pt(9)
p3.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

# Save to destination paths
V8_LOCAL = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_v8.pptx"
V8_DOWNLOADS = r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation_v8.pptx"

prs.save(V8_LOCAL)
print(f"Saved: {V8_LOCAL}")
prs.save(V8_DOWNLOADS)
print(f"Saved: {V8_DOWNLOADS}")

# Also update canonical paths
for path in [
    r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx",
    r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation_updated.pptx",
    r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_updated.pptx",
    r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation_v7.pptx"
]:
    try:
        prs.save(path)
        print(f"Also updated: {path}")
    except Exception as e:
        print(f"Note: {path} is locked: {e}")

print("Thank You slide successfully rebuilt!")
