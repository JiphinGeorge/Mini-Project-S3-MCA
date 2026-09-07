import pptx
from pptx.util import Inches, Pt
import subprocess, time

PPTX_PATH = r"D:\Antigravity Projects\Mini Project S3 MCA\Medical_Specialty_Classification_Presentation.pptx"
prs = pptx.Presentation(PPTX_PATH)

s16 = prs.slides[15] # Slide 16

for sh in s16.shapes:
    if sh.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
        # Scale to fit height = 4.75 in
        # Original ratio: w/h ~ 1.256
        target_h = Inches(4.75)
        aspect_ratio = sh.width / sh.height
        target_w = target_h * aspect_ratio
        sh.height = target_h
        sh.width = int(target_w)
        # Center horizontally on slide
        sh.left = int((Inches(13.333) - target_w) / 2)
        sh.top = Inches(2.05)
        print(f"Fixed Slide 16 Picture: w={sh.width/914400:.2f} in, h={sh.height/914400:.2f} in, left={sh.left/914400:.2f} in, top={sh.top/914400:.2f} in")

# Stop PowerPoint if open
subprocess.run(["powershell", "-Command", "Stop-Process -Name POWERPNT -Force -ErrorAction SilentlyContinue"], capture_output=True)
time.sleep(1)

prs.save(PPTX_PATH)
prs.save(r"C:\Users\jiphi\Downloads\Medical_Specialty_Classification_Presentation.pptx")
print("Slide 16 image scaled and centered successfully!")
